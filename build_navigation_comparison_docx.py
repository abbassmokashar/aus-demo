from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"C:\Users\MC\Desktop\test\AUS Navigation Comparison.docx"
NAVY = "22295F"
LIGHT_BLUE = "EEF6FD"
PALE = "F7F9FC"
GRAY = "D9D9D9"
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)

def set_cell_border(cell, color=GRAY):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '6')
        element.set(qn('w:color'), color)

def set_cell_margin(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn('w:' + side))
        if node is None:
            node = OxmlElement('w:' + side)
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')

def write_cell(cell, text, bold=False, color=BLACK, size=9.2):
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.06
    run = p.add_run(text)
    run.bold = bold
    run.font.name = 'Aptos'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    run.font.size = Pt(size)
    run.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margin(cell)
    set_cell_border(cell)

def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, head in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Inches(widths[i])
        set_cell_shading(cell, NAVY)
        write_cell(cell, head, bold=True, color=WHITE, size=9.1)
    for r_i, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].width = Inches(widths[i])
            if r_i % 2 == 1:
                set_cell_shading(cells[i], PALE)
            write_cell(cells[i], value)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    r.font.color.rgb = BLACK
    return p

def add_bullets(doc, entries):
    for entry in entries:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.1
        p.add_run(entry)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.68)
section.right_margin = Inches(0.68)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].font.color.rgb = BLACK
for name, size in [('Title', 24), ('Heading 1', 15), ('Heading 2', 11.5)]:
    st = styles[name]
    st.font.name = 'Aptos Display' if name != 'Normal' else 'Aptos'
    st._element.rPr.rFonts.set(qn('w:ascii'), st.font.name)
    st._element.rPr.rFonts.set(qn('w:hAnsi'), st.font.name)
    st.font.size = Pt(size)
    st.font.color.rgb = BLACK
    st.font.bold = True

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(5)
title.add_run('AUS Navigation Comparison')
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(14)
subtitle.paragraph_format.line_spacing = 1.15
run = subtitle.add_run('Required navigation compared with the current website implementation')
run.italic = True
run.font.size = Pt(11)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
p.add_run('Conclusion. ').bold = True
p.add_run('The navigation needs a structural rebuild. The current desktop header exposes only Programs, About, and Campus Life, while the directive requires six top-level sections and four global utility actions. The shared menu pattern is currently used across all 37 HTML pages in the website folder.')

add_heading(doc, 'Top Level Navigation', 1)
add_table(doc,
    ['Required section', 'Current implementation', 'Required action'],
    [
        ('Programs', 'Present in desktop header and off canvas menu. Includes degree groups and individual program links.', 'Keep and expand. Add Find Your Program and Compare Programs. Rename Bachelors Degree and Masters Degree.'),
        ('Admissions', 'Only in the off canvas menu. Contains Admissions and Financing, Cost Calculator, and Academic Calendar.', 'Promote to desktop header and rebuild with the complete requested submenu.'),
        ('Tuition and Scholarships', 'Missing. Cost Calculator sits under Admissions.', 'Add as a top level section. Move Cost Calculator here and add finance links.'),
        ('Careers and Outcomes', 'Missing.', 'Add as a top level section with the full required submenu.'),
        ('Student Life', 'Present as Campus Life with only Campus Life and Student Life links.', 'Rename and expand to the requested Student Life structure.'),
        ('About AUS', 'Present as About with a partial list of institutional links.', 'Rename and reorganize to match the directive.'),
    ], [1.18, 2.58, 2.58])

add_heading(doc, 'Programs', 1)
add_table(doc, ['Directive item', 'Current status', 'Action'], [
    ('Find Your Program', 'Missing', 'Add and link to the guided Program Finder.'),
    ('Bachelors Programs', 'Present as Bachelors Degree with program links', 'Rename and retain.'),
    ('Masters Programs', 'Present as Masters Degree with program links', 'Rename and retain.'),
    ('Doctorate', 'Present', 'Retain.'),
    ('Swiss Federal Diploma', 'Present', 'Retain.'),
    ('Compare Programs', 'Missing', 'Add and link to program comparison.'),
], [1.75, 2.25, 2.34])

add_heading(doc, 'Admissions', 1)
add_table(doc, ['Directive item', 'Current status', 'Action'], [
    ('How to Apply', 'Missing', 'Add.'),
    ('Entry Requirements', 'Missing', 'Add.'),
    ('Required Documents', 'Missing', 'Add.'),
    ('Application Deadlines', 'Missing', 'Add; place Academic Calendar here if retained.'),
    ('English Requirements', 'Missing', 'Add.'),
    ('Visa and Immigration', 'Missing', 'Add.'),
    ('Transfer Students', 'Missing', 'Add.'),
    ('FAQ', 'Not under Admissions; currently placed under About', 'Add here and relocate or reuse the current FAQ appropriately.'),
], [1.75, 2.25, 2.34])

add_heading(doc, 'Tuition Scholarships and Careers', 1)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(5)
p.add_run('Tuition and Scholarships. ').bold = True
p.add_run('This is a completely new top-level section. Add Tuition Fees, Scholarships, Financial Aid, Payment Options, Living Costs, Cost Calculator, and FAQ. Move Cost Calculator out of Admissions.')
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(5)
p.add_run('Careers and Outcomes. ').bold = True
p.add_run('This is also a completely new top-level section. Add Career Services, Graduate Outcomes, Internships, Employers, Industry Connections, Alumni, and Student Success Stories.')

add_heading(doc, 'Student Life', 1)
add_table(doc, ['Directive item', 'Current status', 'Action'], [
    ('Campus', 'Present as Campus Life', 'Rename and position as Campus.'),
    ('Housing', 'Missing', 'Add.'),
    ('Living in Switzerland', 'Missing', 'Add.'),
    ('Student Activities and Clubs', 'Missing', 'Add both links.'),
    ('International Community', 'Missing', 'Add.'),
    ('Student Support', 'Missing', 'Add.'),
    ('Transportation', 'Missing', 'Add.'),
    ('Cost of Living', 'Missing', 'Add.'),
    ('Campus Visits', 'Missing', 'Add.'),
], [1.75, 2.25, 2.34])

add_heading(doc, 'About AUS', 1)
add_table(doc, ['Directive item', 'Current status', 'Action'], [
    ('About AUS, History, Faculty, Governance', 'Present', 'Retain.'),
    ('Accreditation and Recognition', 'Present as Accreditations and Memberships', 'Rename and separate accreditation, recognition, rankings, and partnerships.'),
    ('Rankings and Leadership', 'Missing', 'Add. Leadership may be separated from Governance after a content decision.'),
    ('Academic Partners', 'Only Tiffin collaboration exists as an individual link', 'Add the parent item and nest or link Tiffin appropriately.'),
    ('Industry Partners', 'Only IBM collaboration exists as an individual link', 'Add the parent item and nest or link IBM appropriately.'),
    ('News', 'Missing', 'Add.'),
    ('Contact', 'Global Contact action exists', 'Retain globally; include within About only if needed by the final information architecture.'),
], [1.75, 2.25, 2.34])

add_heading(doc, 'Global Utility Actions', 1)
add_table(doc, ['Directive requirement', 'Current status', 'Action'], [
    ('Search or Ask AUS', 'Search input exists with Ask AUS Anything placeholder', 'Keep. Label consistently as Search or Ask AUS and implement grouped search results.'),
    ('My Programs', 'Missing', 'Add and connect to saved programs.'),
    ('Talk to Admissions', 'Present as Contact and links to a HubSpot form', 'Rename and confirm it opens the admissions specific contact flow.'),
    ('Apply', 'Present', 'Keep as the strongest conversion action. Standardize visible wording to Apply or Apply Now.'),
], [1.75, 2.25, 2.34])

add_heading(doc, 'Required Relocation and Cleanup', 1)
add_bullets(doc, [
    'Rename About to About AUS and Campus Life to Student Life.',
    'Split Admissions and Financing: the admissions hub stays under Admissions, while tuition and finance content moves to Tuition and Scholarships.',
    'Move Cost Calculator from Admissions to Tuition and Scholarships.',
    'Move the generic FAQ from About to Admissions and Tuition and Scholarships, using either topic anchors or dedicated section FAQs.',
    'Move Academic Calendar beneath Application Deadlines or a relevant Admissions resource location.',
    'Place the IBM and Tiffin collaboration pages beneath Industry Partners and Academic Partners respectively.',
    'Keep Policies and Procedures in the footer rather than the main navigation.',
    'Remove the hidden Accreditations TEMP and Memberships placeholder link.',
])

add_heading(doc, 'Implementation Requirements', 1)
add_bullets(doc, [
    'Desktop and mobile must expose the same six primary sections and the same essential paths.',
    'Use a desktop mega menu and a mobile accordion if desired, but keep their information architecture identical.',
    'Each primary navigation item must be directly understandable and readily navigable on desktop and mobile.',
    'Do not publish a new navigation destination with a placeholder link.',
    'Apply remains visually stronger than Talk to Admissions.',
])

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run('AUS Navigation Comparison')
footer_run.font.name = 'Aptos'
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor(90, 90, 90)

doc.save(OUT)
print(OUT)
